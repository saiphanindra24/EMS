import csv
import io
from datetime import date, datetime, timedelta
from decimal import Decimal
from django.http import HttpResponse
from django.utils import timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def resolve_date_period(params, default_period="this_month"):
    """
    Parses and validates reporting period and date range from query parameters.
    Supports presets: today, this_week, this_month, last_month, this_quarter, this_year, custom.
    Returns: (start_date, end_date, period_label)
    Raises ValueError with descriptive message if dates are invalid or start_date > end_date.
    """
    period = params.get("period", default_period).lower().strip()
    today = timezone.localdate()

    if period == "today":
        start_date = today
        end_date = today
        label = f"Today ({today.isoformat()})"
    elif period == "this_week":
        # Monday to Sunday of current week
        start_date = today - timedelta(days=today.weekday())
        end_date = start_date + timedelta(days=6)
        label = f"This Week ({start_date.isoformat()} to {end_date.isoformat()})"
    elif period == "this_month":
        start_date = today.replace(day=1)
        # Next month first day minus one day
        next_month = (start_date.replace(day=28) + timedelta(days=4)).replace(day=1)
        end_date = next_month - timedelta(days=1)
        label = f"This Month ({start_date.strftime('%B %Y')})"
    elif period == "last_month":
        first_of_this_month = today.replace(day=1)
        end_date = first_of_this_month - timedelta(days=1)
        start_date = end_date.replace(day=1)
        label = f"Last Month ({start_date.strftime('%B %Y')})"
    elif period == "this_quarter":
        quarter = (today.month - 1) // 3 + 1
        start_month = 3 * quarter - 2
        start_date = date(today.year, start_month, 1)
        end_month = start_month + 2
        next_month = (date(today.year, end_month, 28) + timedelta(days=4)).replace(day=1)
        end_date = next_month - timedelta(days=1)
        label = f"Q{quarter} {today.year} ({start_date.isoformat()} to {end_date.isoformat()})"
    elif period == "this_year":
        start_date = date(today.year, 1, 1)
        end_date = date(today.year, 12, 31)
        label = f"Year {today.year}"
    elif period == "custom" or ("start_date" in params or "end_date" in params):
        start_str = params.get("start_date")
        end_str = params.get("end_date")

        if not start_str or not end_str:
            raise ValueError("Both 'start_date' and 'end_date' (YYYY-MM-DD) are required for custom period.")

        try:
            start_date = datetime.strptime(start_str.strip(), "%Y-%m-%d").date()
        except ValueError:
            raise ValueError("Invalid start_date format. Expected YYYY-MM-DD.")

        try:
            end_date = datetime.strptime(end_str.strip(), "%Y-%m-%d").date()
        except ValueError:
            raise ValueError("Invalid end_date format. Expected YYYY-MM-DD.")

        if start_date > end_date:
            raise ValueError("start_date cannot be later than end_date.")

        label = f"Custom Range ({start_date.isoformat()} to {end_date.isoformat()})"
    else:
        # Fallback to this month
        start_date = today.replace(day=1)
        next_month = (start_date.replace(day=28) + timedelta(days=4)).replace(day=1)
        end_date = next_month - timedelta(days=1)
        label = f"This Month ({start_date.strftime('%B %Y')})"

    return start_date, end_date, label


def generate_csv_response(filename_base, headers, rows, metadata=None):
    """
    Generates a streaming/clean CSV HttpResponse with utf-8-sig encoding for full Excel compatibility.
    Includes optional metadata header rows at the top.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    if metadata:
        for k, v in metadata.items():
            writer.writerow([f"# {k}", str(v)])
        writer.writerow([])  # Blank separator line

    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)

    csv_content = output.getvalue()
    output.close()

    response = HttpResponse(csv_content, content_type="text/csv; charset=utf-8-sig")
    safe_filename = f"{filename_base}_{timezone.localdate().isoformat()}.csv"
    response["Content-Disposition"] = f'attachment; filename="{safe_filename}"'
    return response


def generate_excel_response(filename_base, sheet_title, headers, rows, metadata=None):
    """
    Generates a beautifully styled .xlsx HttpResponse using openpyxl.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel limits sheet names to 31 chars

    # Palette
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    meta_title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    meta_font = Font(name="Calibri", size=10, italic=True, color="475569")
    data_font = Font(name="Calibri", size=10, color="1E293B")
    border_side = Side(style="thin", color="E2E8F0")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    current_row = 1

    if metadata:
        ws.cell(row=current_row, column=1, value="EMWTS Workplace Management System").font = meta_title_font
        current_row += 1
        for k, v in metadata.items():
            cell = ws.cell(row=current_row, column=1, value=f"{k}: {v}")
            cell.font = meta_font
            current_row += 1
        current_row += 1  # Blank spacing

    # Table Header
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if "ID" in header or "Date" in header or "Rate" in header or "%" in header else "left", vertical="center")
        cell.border = cell_border
    ws.row_dimensions[current_row].height = 26
    current_row += 1

    # Data rows
    for r_idx, row_data in enumerate(rows):
        is_even = (r_idx % 2 == 1)
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = data_font
            cell.border = cell_border
            if is_even:
                cell.fill = alt_fill

            # Alignment logic
            if isinstance(value, (int, float, Decimal)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(value, (date, datetime)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Auto column width adjustment
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip title row from calculation
            if cell.row < (len(metadata) + 2 if metadata else 1):
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    safe_filename = f"{filename_base}_{timezone.localdate().isoformat()}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{safe_filename}"'
    return response


PRODUCTIVITY_DEFINITIONS = {
    "task_completion_rate": "Percentage of tasks assigned that reached COMPLETED status within the reporting window.",
    "on_time_delivery_rate": "Percentage of completed tasks delivered on or prior to their assigned due dates.",
    "estimation_efficiency": "Ratio comparing planned estimated hours to actual logged hours. 100% represents optimal alignment; >100% indicates under-budget velocity.",
    "attendance_reliability": "Ratio of attended work hours against the standard schedule (8h/day baseline), accounting for approved leaves.",
    "composite_productivity_score": "Multifaceted performance index calculated as: 40% Task Completion + 25% On-Time Delivery + 25% Attendance Reliability + 10% Estimation Efficiency. Provides a balanced view rather than crude task volume.",
}
