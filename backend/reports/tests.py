import io
from datetime import date, timedelta
from decimal import Decimal
from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone
import openpyxl
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import Role
from attendance.models import AttendanceRecord, AttendanceStatus
from employees.models import Department, Employee, EmploymentStatus, EmploymentType
from leave_management.models import LeaveRequest, LeaveStatus, LeaveType
from tasks.models import Task, TaskPriority, TaskStatus

User = get_user_model()


class ReportAPITests(APITestCase):
    def setUp(self):
        # Departments
        self.dept_eng = Department.objects.create(name="Engineering", code="ENG")
        self.dept_sales = Department.objects.create(name="Sales", code="SLS")

        # Super Admin
        self.admin = User.objects.create_superuser(
            email="admin@emwts.local",
            password="AdminPassword2026!",
            first_name="Admin",
            last_name="Boss",
        )
        self.admin_profile = Employee.objects.create(
            user=self.admin,
            employee_id="EMP-ADM-001",
            department=self.dept_eng,
            base_salary=Decimal("120000.00"),
            national_id="SSN-ADMIN-01",
        )

        # HR Admin
        self.hr = User.objects.create_user(
            email="hr@emwts.local",
            password="HrPassword2026!",
            first_name="Helen",
            last_name="HR",
            role=Role.HR_ADMIN,
        )
        self.hr_profile = Employee.objects.create(
            user=self.hr,
            employee_id="EMP-HR-001",
            department=self.dept_eng,
            base_salary=Decimal("95000.00"),
            national_id="SSN-HR-01",
        )

        # Manager
        self.manager = User.objects.create_user(
            email="manager@emwts.local",
            password="ManagerPassword2026!",
            first_name="Manoj",
            last_name="Manager",
            role=Role.MANAGER,
        )
        self.manager_profile = Employee.objects.create(
            user=self.manager,
            employee_id="EMP-MGR-001",
            department=self.dept_eng,
            base_salary=Decimal("80000.00"),
            national_id="SSN-MGR-01",
        )

        # Employee 1 (reports to manager)
        self.emp1 = User.objects.create_user(
            email="emp1@emwts.local",
            password="EmpPassword2026!",
            first_name="Esha",
            last_name="Engineer",
            role=Role.EMPLOYEE,
        )
        self.emp1_profile = Employee.objects.create(
            user=self.emp1,
            employee_id="EMP-DEV-001",
            department=self.dept_eng,
            manager=self.manager_profile,
            base_salary=Decimal("60000.00"),
            national_id="SSN-EMP-01",
            bank_account_number="ACC-123456",
            address="123 Tech Lane",
            date_of_birth=date(1995, 5, 15),
        )

        # Employee 2 (Sales department, not managed by engineering manager)
        self.emp2 = User.objects.create_user(
            email="emp2@emwts.local",
            password="EmpPassword2026!",
            first_name="Sam",
            last_name="Sales",
            role=Role.EMPLOYEE,
        )
        self.emp2_profile = Employee.objects.create(
            user=self.emp2,
            employee_id="EMP-SLS-001",
            department=self.dept_sales,
            base_salary=Decimal("55000.00"),
            national_id="SSN-EMP-02",
        )

        # Leave Type
        self.vacation_type = LeaveType.objects.create(
            name="Vacation Leave",
            code="VAC",
            is_paid=True,
        )

        # Seed test date range: last 5 days
        self.today = timezone.localdate()
        self.d1 = self.today - timedelta(days=4)
        self.d2 = self.today - timedelta(days=3)
        self.d3 = self.today - timedelta(days=2)
        self.d4 = self.today - timedelta(days=1)
        self.d5 = self.today

        # Seed attendance records for emp1
        AttendanceRecord.objects.create(
            employee=self.emp1_profile,
            date=self.d1,
            status=AttendanceStatus.PRESENT,
            work_duration_hours=Decimal("8.00"),
            work_duration_minutes=480,
        )
        AttendanceRecord.objects.create(
            employee=self.emp1_profile,
            date=self.d2,
            status=AttendanceStatus.LATE,
            work_duration_hours=Decimal("7.50"),
            work_duration_minutes=450,
        )
        AttendanceRecord.objects.create(
            employee=self.emp1_profile,
            date=self.d3,
            status=AttendanceStatus.PRESENT,
            work_duration_hours=Decimal("8.00"),
            work_duration_minutes=480,
        )

        # Seed attendance for emp2
        AttendanceRecord.objects.create(
            employee=self.emp2_profile,
            date=self.d1,
            status=AttendanceStatus.PRESENT,
            work_duration_hours=Decimal("8.00"),
            work_duration_minutes=480,
        )

        # Seed tasks for emp1
        self.task1 = Task.objects.create(
            task_id="TSK-001",
            title="Implement Core API",
            assignee=self.emp1_profile,
            assigning_manager=self.manager,
            department=self.dept_eng,
            priority=TaskPriority.HIGH,
            status=TaskStatus.COMPLETED,
            progress=100,
            estimated_hours=Decimal("10.00"),
            actual_hours=Decimal("9.00"),
            due_date=self.d4,
        )
        self.task2 = Task.objects.create(
            task_id="TSK-002",
            title="Write Tests",
            assignee=self.emp1_profile,
            assigning_manager=self.manager,
            department=self.dept_eng,
            priority=TaskPriority.MEDIUM,
            status=TaskStatus.IN_PROGRESS,
            progress=50,
            estimated_hours=Decimal("8.00"),
            actual_hours=Decimal("4.00"),
            due_date=self.today + timedelta(days=2),
        )

        # Seed task for emp2
        self.task_sales = Task.objects.create(
            task_id="TSK-003",
            title="Sales Pitch Prep",
            assignee=self.emp2_profile,
            department=self.dept_sales,
            priority=TaskPriority.LOW,
            status=TaskStatus.COMPLETED,
            progress=100,
            estimated_hours=Decimal("5.00"),
            actual_hours=Decimal("5.00"),
            due_date=self.d2,
        )

        # Seed leave for emp1
        self.leave1 = LeaveRequest.objects.create(
            employee=self.emp1_profile,
            leave_type=self.vacation_type,
            start_date=self.d4,
            end_date=self.d4,
            duration_days=Decimal("1.0"),
            reason="Personal travel",
            status=LeaveStatus.APPROVED,
        )

    # --------------------------------------------------------------------------
    # 1. OVERVIEW ENDPOINT
    # --------------------------------------------------------------------------
    def test_overview_endpoint(self):
        self.client.force_authenticate(user=self.admin)
        res = self.client.get(reverse("reports:overview"))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("available_reports", res.data)
        self.assertIn("high_level_stats", res.data)
        self.assertIn("definitions", res.data)
        self.assertEqual(len(res.data["available_reports"]), 6)

    # --------------------------------------------------------------------------
    # 2. EMPLOYEE REPORTS & PERMISSIONS (Sensitive Data Masking)
    # --------------------------------------------------------------------------
    def test_employee_report_admin_sees_sensitive_fields(self):
        self.client.force_authenticate(user=self.admin)
        res = self.client.get(reverse("reports:employees"))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(res.data["count"], 5)

        # Admin should see sensitive_data
        emp1_data = next((e for e in res.data["results"] if e["employee_id"] == "EMP-DEV-001"), None)
        self.assertIsNotNone(emp1_data)
        self.assertIn("sensitive_data", emp1_data)
        self.assertEqual(emp1_data["sensitive_data"]["national_id"], "SSN-EMP-01")
        self.assertEqual(emp1_data["sensitive_data"]["bank_account_number"], "ACC-123456")

    def test_employee_report_manager_cannot_see_sensitive_fields(self):
        self.client.force_authenticate(user=self.manager)
        res = self.client.get(reverse("reports:employees"))
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        # Manager should only see their scoped employees (self, emp1, admin in dept), NOT emp2 from Sales
        emp_ids = [e["employee_id"] for e in res.data["results"]]
        self.assertIn("EMP-DEV-001", emp_ids)
        self.assertNotIn("EMP-SLS-001", emp_ids)

        # Crucial: Manager should NOT see sensitive personal financial data
        emp1_data = next((e for e in res.data["results"] if e["employee_id"] == "EMP-DEV-001"), None)
        self.assertIsNotNone(emp1_data)
        self.assertNotIn("sensitive_data", emp1_data)

    def test_employee_report_employee_only_sees_self(self):
        self.client.force_authenticate(user=self.emp1)
        res = self.client.get(reverse("reports:employees"))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["count"], 1)
        self.assertEqual(res.data["results"][0]["employee_id"], "EMP-DEV-001")

    # --------------------------------------------------------------------------
    # 3. ATTENDANCE REPORTS & DATE FILTERS
    # --------------------------------------------------------------------------
    def test_attendance_report_date_filtering(self):
        self.client.force_authenticate(user=self.admin)
        url = reverse("reports:attendance")

        # Custom date range including d1 and d2 only
        res = self.client.get(url, {
            "period": "custom",
            "start_date": self.d1.isoformat(),
            "end_date": self.d2.isoformat(),
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        # Should contain d1 (emp1 & emp2) and d2 (emp1) = 3 records
        self.assertEqual(res.data["count"], 3)
        self.assertEqual(res.data["summary"]["present_count"], 2)
        self.assertEqual(res.data["summary"]["late_count"], 1)

    def test_attendance_report_invalid_date_boundary(self):
        self.client.force_authenticate(user=self.admin)
        url = reverse("reports:attendance")
        # start_date > end_date
        res = self.client.get(url, {
            "period": "custom",
            "start_date": "2026-10-10",
            "end_date": "2026-10-01",
        })
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("error", res.data)

    def test_attendance_report_manager_scoping(self):
        self.client.force_authenticate(user=self.manager)
        url = reverse("reports:attendance")
        res = self.client.get(url, {
            "period": "custom",
            "start_date": self.d1.isoformat(),
            "end_date": self.d5.isoformat(),
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        # Manager only sees emp1 records (3 records), not emp2 (Sales)
        for r in res.data["results"]:
            self.assertNotEqual(r["employee"]["employee_id"], "EMP-SLS-001")

    # --------------------------------------------------------------------------
    # 4. TASK REPORTS & FILTERS
    # --------------------------------------------------------------------------
    def test_task_report_status_and_priority_filter(self):
        self.client.force_authenticate(user=self.admin)
        url = reverse("reports:tasks")

        res_completed = self.client.get(url, {"status": "COMPLETED"})
        self.assertEqual(res_completed.status_code, status.HTTP_200_OK)
        for t in res_completed.data["results"]:
            self.assertEqual(t["status"], "COMPLETED")

        res_priority = self.client.get(url, {"priority": "HIGH"})
        self.assertEqual(res_priority.status_code, status.HTTP_200_OK)
        for t in res_priority.data["results"]:
            self.assertEqual(t["priority"], "HIGH")

    # --------------------------------------------------------------------------
    # 5. LEAVE REPORTS
    # --------------------------------------------------------------------------
    def test_leave_report_filtering(self):
        self.client.force_authenticate(user=self.admin)
        url = reverse("reports:leaves")

        res = self.client.get(url, {
            "status": "APPROVED",
            "period": "custom",
            "start_date": self.d1.isoformat(),
            "end_date": self.d5.isoformat(),
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["count"], 1)
        self.assertEqual(res.data["summary"]["total_approved_days"], 1.0)

    # --------------------------------------------------------------------------
    # 6. DEPARTMENT REPORTS
    # --------------------------------------------------------------------------
    def test_department_report(self):
        self.client.force_authenticate(user=self.admin)
        url = reverse("reports:departments")

        res = self.client.get(url)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(res.data["count"], 2)

        eng_dept = next((d for d in res.data["results"] if d["code"] == "ENG"), None)
        self.assertIsNotNone(eng_dept)
        self.assertGreaterEqual(eng_dept["active_employees"], 3)
        self.assertGreaterEqual(eng_dept["total_tasks"], 2)

    # --------------------------------------------------------------------------
    # 7. PRODUCTIVITY REPORTS (Multi-Factor Index & Definitions)
    # --------------------------------------------------------------------------
    def test_productivity_report_calculations_and_definitions(self):
        self.client.force_authenticate(user=self.admin)
        url = reverse("reports:productivity")

        res = self.client.get(url, {
            "period": "custom",
            "start_date": self.d1.isoformat(),
            "end_date": self.d5.isoformat(),
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("definitions", res.data)
        self.assertIn("task_completion_rate", res.data["definitions"])
        self.assertIn("composite_productivity_score", res.data["definitions"])

        # Check emp1 metrics
        emp1_prod = next((p for p in res.data["results"] if p["employee_id"] == "EMP-DEV-001"), None)
        self.assertIsNotNone(emp1_prod)
        self.assertEqual(emp1_prod["total_tasks"], 2)
        self.assertEqual(emp1_prod["completed_tasks"], 1)
        self.assertEqual(emp1_prod["task_completion_rate"], 50.0)
        self.assertGreater(emp1_prod["total_logged_hours"], 0)
        self.assertGreater(emp1_prod["composite_productivity_score"], 0)
        self.assertIn(emp1_prod["performance_tier"], ["Exceptional", "High Performing", "Satisfactory", "Needs Attention"])

    # --------------------------------------------------------------------------
    # 8. EXPORT CSV & EXCEL (.XLSX) WITH AUTH CONTROLS
    # --------------------------------------------------------------------------
    def test_export_csv_admin_vs_manager(self):
        # Admin gets CSV with sensitive columns
        self.client.force_authenticate(user=self.admin)
        res_admin = self.client.get(reverse("reports:employees"), {"export": "csv"})
        self.assertEqual(res_admin.status_code, status.HTTP_200_OK)
        self.assertEqual(res_admin["Content-Type"], "text/csv; charset=utf-8-sig")
        content_admin = res_admin.content.decode("utf-8-sig")
        self.assertIn("Base Salary", content_admin)
        self.assertIn("National ID", content_admin)
        self.assertIn("SSN-EMP-01", content_admin)

        # Manager gets CSV WITHOUT sensitive columns
        self.client.force_authenticate(user=self.manager)
        res_mgr = self.client.get(reverse("reports:employees"), {"export": "csv"})
        self.assertEqual(res_mgr.status_code, status.HTTP_200_OK)
        content_mgr = res_mgr.content.decode("utf-8-sig")
        self.assertNotIn("Base Salary", content_mgr)
        self.assertNotIn("National ID", content_mgr)
        self.assertNotIn("SSN-EMP-01", content_mgr)

    def test_export_excel_workbook_structure(self):
        self.client.force_authenticate(user=self.admin)
        url = reverse("reports:productivity")

        res = self.client.get(url, {
            "export": "xlsx",
            "period": "custom",
            "start_date": self.d1.isoformat(),
            "end_date": self.d5.isoformat(),
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheetml.sheet", res["Content-Type"])

        # Verify binary stream with openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertIn("Productivity", wb.sheetnames)
        ws = wb["Productivity"]
        # Check that rows and columns were created
        self.assertGreater(ws.max_row, 3)
        self.assertGreater(ws.max_column, 5)
